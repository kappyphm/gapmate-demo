"""Stage 1 — extract text and layout facts from an uploaded file.

Also validates the upload (type/size/emptiness) and runs an OCR fallback (Tesseract vie+eng)
for scanned PDFs. The ``layout`` dict feeds the ATS Group A rules (docs/04 Part C).
"""

import io

MAX_BYTES = 5 * 1024 * 1024
SUPPORTED_EXT = {"pdf", "docx", "txt", "md"}


class ExtractError(ValueError):
    """Raised when a document cannot be safely extracted."""


def extract_text(data: bytes, filename: str) -> tuple[str, dict]:
    if not data:
        raise ExtractError("Empty file.")
    if len(data) > MAX_BYTES:
        raise ExtractError("File exceeds the 5 MB limit.")

    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext not in SUPPORTED_EXT:
        raise ExtractError(f"Unsupported file type: .{ext or 'unknown'}")

    if ext == "pdf":
        return _extract_pdf(data)
    if ext == "docx":
        return _extract_docx(data)
    return _extract_plain(data, ext)


def _base_layout(file_type: str) -> dict:
    return {
        "file_type": file_type,
        "pages": 1,
        "has_text": False,
        "used_ocr": False,
        "single_column": True,
        "has_tables": False,
        "header_footer_content": False,
        "standard_fonts": True,
        "has_text_boxes": False,
        "blocks": [],
    }


def _extract_plain(data: bytes, ext: str) -> tuple[str, dict]:
    text = data.decode("utf-8", errors="ignore")
    layout = _base_layout(ext)
    layout["has_text"] = bool(text.strip())
    return text, layout


def _extract_docx(data: bytes) -> tuple[str, dict]:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractError("Corrupt or unreadable DOCX file.") from exc

    text = "\n".join(p.text for p in document.paragraphs)
    layout = _base_layout("docx")
    layout["has_text"] = bool(text.strip())
    layout["has_tables"] = len(document.tables) > 0
    return text, layout


def _extract_pdf(data: bytes) -> tuple[str, dict]:
    import fitz  # PyMuPDF

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:  # noqa: BLE001
        raise ExtractError("Corrupt or unreadable PDF file.") from exc

    if doc.needs_pass:
        doc.close()
        raise ExtractError("Password-protected PDF.")

    parts: list[str] = []
    blocks: list[dict] = []
    fonts: set[str] = set()
    has_tables = False
    header_footer = False

    for page in doc:
        parts.append(page.get_text("text"))
        rect = page.rect
        width = rect.width or 1
        height = rect.height or 1
        page_blocks = page.get_text("blocks")
        blocks.extend({"page": page.number, "bbox": list(b[:4]), "text": b[4]} for b in page_blocks)
        has_tables = has_tables or _has_tables(page)
        fonts.update(_fonts(page))
        header_footer = header_footer or _content_in_margins(page_blocks, width, height)

    text = "\n".join(parts)
    used_ocr = False
    if not text.strip():
        text = _ocr_pdf(doc)
        used_ocr = True

    layout = _base_layout("pdf")
    layout.update(
        pages=doc.page_count,
        has_text=bool(text.strip()),
        used_ocr=used_ocr,
        single_column=_is_single_column(blocks, doc),
        has_tables=has_tables,
        header_footer_content=header_footer,
        standard_fonts=len(fonts) <= 4,
    )
    doc.close()
    return text, layout


def _has_tables(page) -> bool:
    try:
        return len(page.find_tables().tables) > 0
    except Exception:  # noqa: BLE001
        return False


def _fonts(page) -> set[str]:
    try:
        return {font[3].split("+")[-1] for font in page.get_fonts()}
    except Exception:  # noqa: BLE001
        return set()


def _content_in_margins(blocks: list, width: float, height: float) -> bool:
    import re

    contact = re.compile(r"@|\+?\d[\d\s.\-]{7,}\d")
    for b in blocks:
        text = (b[4] or "").strip()
        if not text or len(text) > 80 or not contact.search(text):
            continue
        y0, y1 = b[1], b[3]
        if y1 < 0.12 * height or y0 > 0.88 * height:
            return True
    return False


def _is_single_column(blocks: list, doc) -> bool:
    multi = False
    for page_number in range(doc.page_count):
        page = doc[page_number]
        width = page.rect.width or 1
        xs = [b["bbox"][0] for b in blocks if b["page"] == page_number and b["text"].strip()]
        if len(xs) < 4:
            continue
        right = [x for x in xs if x > 0.45 * width]
        left = [x for x in xs if x <= 0.45 * width]
        if len(right) >= 2 and len(left) >= 2:
            multi = True
    return not multi


def _ocr_pdf(doc) -> str:
    try:
        import pytesseract
        from PIL import Image
    except Exception:  # noqa: BLE001
        return ""

    texts: list[str] = []
    for page in doc:
        try:
            pix = page.get_pixmap(dpi=200)
            image = Image.open(io.BytesIO(pix.tobytes("png")))
        except Exception:  # noqa: BLE001
            continue
        try:
            texts.append(pytesseract.image_to_string(image, lang="vie+eng"))
        except Exception:  # noqa: BLE001
            texts.append(pytesseract.image_to_string(image))
    return "\n".join(texts)
